from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


OUT = Path(__file__).parent


def configure(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)


explicit = Document()
configure(explicit)
for page_number in range(1, 4):
    explicit.add_heading(f"Pagination fixture - page {page_number}", level=1)
    for line in range(1, 10):
        explicit.add_paragraph(
            f"Page {page_number}, paragraph {line}. This sentence verifies that "
            "Word page boundaries remain separate in the browser."
        )
    if page_number < 3:
        explicit.add_page_break()
explicit.save(OUT / "word-explicit-pages.docx")


flowing = Document()
configure(flowing)
flowing.add_heading("Pagination fixture - flowing content", level=1)
for line in range(1, 125):
    flowing.add_paragraph(
        f"Flow paragraph {line}. A long Word document without explicit page breaks "
        "must still become multiple writable pages in the browser."
    )
flowing.save(OUT / "word-flowing-pages.docx")
